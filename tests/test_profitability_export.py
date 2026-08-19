from __future__ import annotations

import unittest

import pandas as pd
from openpyxl import load_workbook

from modules.exports.profitability_export import build_xlsx


class ProfitabilityExportPeriodTests(unittest.TestCase):
    def test_exports_only_selected_invoice_month_range(self):
        df_prof = pd.DataFrame([
            self._row("ZP-1", "2026-01", 100),
            self._row("ZP-2", "2026-02", 200),
            self._row("ZP-3", "2026-03", 300),
        ])
        result = {
            "df_prof": df_prof,
            "df_czasy": None,
            "df_kliki": None,
            "df_farby_pivot": None,
            "machine_cols": [],
            "klient_col": None,
            "df_material": None,
            "df_tektura": None,
            "df_orders": None,
        }

        output = build_xlsx(
            result, {}, {}, {}, 2.0, 60.0, 40.0,
            month_from="2026-02", month_to="2026-03",
        )
        workbook = load_workbook(output, read_only=True, data_only=True)

        profitability = workbook["Profitability"]
        self.assertEqual(profitability.max_row, 3)  # header + two records
        self.assertIn("Podsum. 2026-02", workbook.sheetnames)
        self.assertIn("Podsum. 2026-03", workbook.sheetnames)
        self.assertNotIn("Podsum. 2026-01", workbook.sheetnames)

    @staticmethod
    def _row(order: str, month: str, sales: float) -> dict:
        return {
            "Zlecenie produkcyjne": order,
            "Data faktury": pd.Timestamp(f"{month}-15"),
            "Miesiąc faktury": month,
            "Sales Value": sales,
            "TPM": sales * 0.7,
            "TPM%": 70.0,
            "CM": sales * 0.5,
            "CM%": 50.0,
        }


if __name__ == "__main__":
    unittest.main()
